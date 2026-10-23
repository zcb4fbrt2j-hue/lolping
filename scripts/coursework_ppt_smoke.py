"""Generate and validate an editable Chinese slide; run from the repository root."""
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'examples/coursework-ppt/中文学术汇报_测试页.pptx'


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string('F5F7FA')

    def text(value, x, y, w, h, size, color='183153'):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        box.text_frame.word_wrap = True
        p = box.text_frame.paragraphs[0]
        p.text = value
        p.font.name = 'Noto Sans CJK SC'
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor.from_string(color)
        return box

    text('研究生学术汇报 · 工具验证', .65, .4, 12, .5, 20)
    text('从研究问题到可追溯的证据', .65, 1.15, 12, .85, 36)
    text('以“问题—证据—分析”组织单页论证', .65, 2.1, 12, .6, 26)
    for x, title, body in [
        (.65, '研究问题', '明确研究对象\n限定讨论边界'),
        (4.85, '证据材料', '记录原始来源\n区分事实与推断'),
        (9.05, '分析结论', '解释证据含义\n说明适用限制'),
    ]:
        text(title, x, 3.25, 3.65, .6, 28)
        text(body, x, 4.15, 3.65, 1.25, 24)
    text('资料来源：本项目 coursework-ppt / SKILL.md（Operating Principle、Required Workflow）',
         .65, 6.1, 12, .6, 16)
    text('方法流程示意；不包含实证数据或研究结果。', .65, 6.75, 12, .4, 16)
    slide.notes_slide.notes_text_frame.text = (
        '测试范围：中文文本、可编辑形状与 OOXML 文件结构。\n'
        '来源：https://github.com/Tang1206cc/codex-academic-presentation-skill\n'
        '上游版本：6aed88f1461ed9eec1ceb5f5f0f935b68efabd56\n'
        '这是一页安装冒烟测试，不代表真实研究结论。'
    )
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    with ZipFile(OUTPUT) as archive:
        assert archive.testzip() is None, 'ZIP CRC failure'
        assert '[Content_Types].xml' in archive.namelist()
        for name in archive.namelist():
            if name.endswith(('.xml', '.rels')):
                etree.fromstring(archive.read(name))
    reopened = Presentation(OUTPUT)
    assert len(reopened.slides) == 1
    texts = [s.text for s in reopened.slides[0].shapes if s.has_text_frame]
    assert '从研究问题到可追溯的证据' in texts
    assert len(texts) == 11
    for shape in reopened.slides[0].shapes:
        assert 0 <= shape.left < shape.left + shape.width <= reopened.slide_width
        assert 0 <= shape.top < shape.top + shape.height <= reopened.slide_height
    print(f'PASS: 1 slide, 11 editable Chinese text boxes, ZIP/XML, content and bounds: {OUTPUT}')


if __name__ == '__main__':
    main()

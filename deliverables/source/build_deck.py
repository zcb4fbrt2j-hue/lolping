"""Editable 15-slide academic deck with verified institutional sources.

Run from any directory: python presentations/船缝里的材料智慧/build_deck.py
"""
from pathlib import Path
from math import sin, cos, pi
import json
import re
from zipfile import ZipFile
from PIL import ImageFont, Image
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BASE = Path(__file__).resolve().parent
OUT = BASE / '一条船缝里的材料智慧_课堂汇报.pptx'
FONT = 'Noto Sans CJK SC'
FONTFILE = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
W, H = 13.333333, 7.5
NAVY, INK, BLUE = '071E33', '16344B', '176B9D'
CYAN, GOLD, WHITE = '52CCCF', 'C9A660', 'FFFFFF'
PAPER, MUTED, PALE = 'F2F6F7', 'AAC0CF', 'DCE9EC'
ORANGE, GREY, WOOD = 'F1BC72', '536B7A', 'B08859'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
LAYOUT = []
NOTES = []
TIMES = [20,50,50,120,30,75,75,65,55,80,75,85,85,35,0]
CURRENT = None
SOURCES = {
    1: 'https://www.ihchina.cn/directory_list',
    2: 'https://www.cscp.org.cn/site/content/3294.html',
    3: 'https://www.gov.cn/yaowen/liebiao/202603/content_7062633.htm',
    4: 'https://www.shanghai.gov.cn/nw12344/20260910/5c1780a7432543a389d9fe119d1e00b1.html',
    5: 'https://www.imo.org/en/OurWork/Environment/Pages/Anti-fouling.aspx',
    6: 'https://www.jmu.edu.cn/info/1041/118240.htm',
    7: 'https://www.mdpi.com/1996-1944/18/7/1531',
    8: 'https://www.mdpi.com/1422-0067/15/12/22142',
}


def rgb(c):
    return RGBColor.from_string(c)


def shape(x, y, w, h, color, kind=MSO_SHAPE.RECTANGLE, border=None, rotation=0):
    s = CURRENT.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(color)
    if border:
        s.line.color.rgb = rgb(border)
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.rotation = rotation
    return s


def line(x1, y1, x2, y2, c=CYAN, width=1.5):
    s = CURRENT.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                    Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    s.line.color.rgb = rgb(c)
    s.line.width = Pt(width)
    return s


def polygon(points, c):
    scale = 1000
    builder = CURRENT.shapes.build_freeform(points[0][0]*scale, points[0][1]*scale,
                                           scale=Inches(1)/scale)
    builder.add_line_segments([(x*scale, y*scale) for x, y in points[1:]], close=True)
    s = builder.convert_to_shape()
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(c)
    s.line.fill.background()
    return s


def wrap(value, size, width):
    font = ImageFont.truetype(FONTFILE, int(round(size*4)))
    max_width = width*72*4
    result = []
    for explicit in str(value).split('\n'):
        current = ''
        for char in explicit:
            if current and font.getlength(current+char) > max_width:
                result.append(current)
                current = char
            else:
                current += char
        result.append(current)
    return result


def text(value, x, y, w, h=None, size=20, color=INK, bold=False,
         align=PP_ALIGN.LEFT, check=True):
    lines = wrap(value, size, w-.06)
    needed = len(lines)*size*1.28/72 + .06
    if h is None:
        h = needed
    assert needed <= h+.10, f'Text height: {value!r}, {needed:.2f}>{h:.2f}'
    h = max(h,needed)
    b = CURRENT.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = b.text_frame
    tf.margin_left = tf.margin_right = Inches(.025)
    tf.margin_top = tf.margin_bottom = Inches(.01)
    tf.word_wrap = False
    for i, value_line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = value_line
        p.font.name, p.font.size, p.font.bold = FONT, Pt(size), bold
        p.font.color.rgb = rgb(color)
        p.alignment = align
        p.line_spacing = Pt(size*1.28)
        p.space_before = p.space_after = Pt(0)
        # Explicit East Asian font prevents fallback in Office renderers.
        rpr = p._p.get_or_add_pPr().get_or_add_defRPr()
        east = etree.SubElement(rpr, '{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
        east.set('typeface', FONT)
    LAYOUT.append(dict(slide=len(prs.slides), text=str(value), x=x, y=y, w=w, h=h,
                       needed=needed, size=size, check=check))
    return b


def label(value, x, y, color=BLUE, w=3, size=13):
    return text(value, x, y, w, max(.32,size*1.28/72+.065), size, color, True)


def new(chapter, title, subtitle, source, dark=False, number=None):
    global CURRENT
    CURRENT = prs.slides.add_slide(prs.slide_layouts[6])
    bg, fg = (NAVY, WHITE) if dark else (PAPER, INK)
    CURRENT.background.fill.solid()
    CURRENT.background.fill.fore_color.rgb = rgb(bg)
    shape(.55, .35, .075, .25, GOLD)
    label(chapter, .78, .31, CYAN if dark else BLUE, w=9)
    label('材料 × 海洋', 11.36, .31, MUTED if dark else GREY, w=1.4, size=12)
    text(title, .62, .93, 12.1, .68, 29, fg, True)
    text(subtitle, .65, 1.75, 12, .55, 18, MUTED if dark else GREY)
    line(.65, 6.82, 12.68, 6.82, '315064' if dark else 'CFDDE2', .7)
    text(source, .65, 6.98, 11.35, .28, 10.5, MUTED if dark else GREY)
    text(f'{len(prs.slides):02d} / 15', 12.02, 6.96, .7, .30, 10.5, MUTED if dark else GREY)
    if number:
        for i in range(5):
            shape(.65+i*.35, 6.65, .25, .035, GOLD if i+1==number else ('315064' if dark else 'CFDDE2'))
    return CURRENT


def arrow(x, y, w=.5, color=CYAN):
    return shape(x, y, w, .28, color, MSO_SHAPE.CHEVRON)


def panel(x, y, w, h, c=WHITE):
    return shape(x, y, w, h, c, MSO_SHAPE.ROUNDED_RECTANGLE)


def photo(filename,x,y,w,h):
    path=BASE/'assets'/filename
    iw,ih=Image.open(path).size
    scale=min(w/iw,h/ih)
    pw,ph=iw*scale,ih*scale
    return CURRENT.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),
                                      width=Inches(pw),height=Inches(ph))


def ship(x, y, w, c=BLUE, traditional=False):
    polygon([(x,y+.55),(x+w,y+.55),(x+w*.86,y+1.10),(x+w*.2,y+1.10)],c)
    if traditional:
        line(x+w*.50,y-.6,x+w*.5,y+.55,GOLD,3)
        polygon([(x+w*.51,y-.55),(x+w*.80,y-.28),(x+w*.51,y+.27)],WOOD)
        for a in [.32,.48,.64,.8]:
            line(x+w*a,y+.57,x+w*a,y+.99,PAPER,1.3)
    else:
        shape(x+w*.70,y+.05,w*.16,.5,'6C8797')
        for row in range(2):
            for col in range(4):
                shape(x+w*.23+col*w*.11,y+.06+row*.23,w*.10,.20,CYAN if row else '497795')
    for d in [0,.14,.28]:
        line(x,y+1.2+d,x+w,y+1.2+d,'557B91',.7)


# 1 — Cover
CURRENT = prs.slides.add_slide(prs.slide_layouts[6])
CURRENT.background.fill.solid(); CURRENT.background.fill.fore_color.rgb=rgb(NAVY)
shape(8.72,0,4.61,7.5,'102E45')
for j in range(7):
    line(8.9, .8+j*.92, 13.05, .8+j*.92, '1B4156', .7)
label('材料与化工  /  研究生思政课程小组汇报', .7, .45, CYAN, 10, 14)
label('15 分钟 · 16:9',11.12,.45,MUTED,1.7,12)
text('一条船缝里的\n材料智慧', .68, 1.37, 8.0, 1.88, 44, WHITE, True)
shape(.72,3.55,.62,.055,GOLD)
text('从传统福船艌缝到现代船舶防护材料\n看海洋强国建设之路',.7,3.94,7.6,1.25,25,WHITE)
text('15 分钟 · 五章论证 · 传统工艺 × 材料创新 × 国家战略',.7,5.57,7.5,.75,17,MUTED)
ship(9.13,1.7,3.55,WOOD,True)
ship(9.13,4.17,3.55,CYAN,False)
label('结构与接缝 → 表面与界面',9.02,6.08,MUTED,3.6,12)
text('图为可编辑示意 · 以非遗资料、成果评价与正式政策文件为依据',.7,6.95,11.7,.3,11,MUTED)

# 2 — Structure vs seal
new('第一章  千年匠心——一条船缝里的中国智慧','隔舱管“范围”，艌缝管“通道”',
    '二者解决不同尺度的进水问题，不能把水密隔舱等同于填缝材料。',
    '来源 [1] 非遗网官方名录数据（项目 11847）；图为原理示意，非文物复原',number=1)
panel(.65,2.55,7.15,3.72)
label('船体结构 · 分舱示意',.95,2.85,BLUE,5)
polygon([(1.08,3.54),(7.25,3.54),(6.50,5.25),(1.88,5.25)],'D9BB89')
for x in [2.36,3.67,4.98,6.29]:
    shape(x,3.57,.085,1.42,WOOD)
shape(3.80,4.32,1.10,.66,'6DACC0')
label('局部进水',3.76,5.48,BLUE,1.5,13)
label('结构分隔 ≠ 接缝密封',1.0,5.94,INK,6,17)
label('局部接缝 · 放大示意',8.26,2.68,BLUE,4.4)
shape(8.28,3.36,1.85,1.32,WOOD)
shape(10.42,3.36,1.85,1.32,WOOD)
shape(10.13,3.36,.29,1.32,GOLD)
for yy in [3.55,3.8,4.05,4.3,4.55]:
    line(10.15,yy,10.38,yy+.08,INK,.9)
text('板缝填塞\n降低渗漏通道的连通性',8.26,5.05,4.35,1.0,22)

# 3 — material roles
new('第一章  千年匠心','一条接缝，也是一次材料与工艺的协同设计',
    '先理解各组分承担什么角色，再讨论如何填塞、压实与维护。',
    '来源 [1] 非遗网原文确认三种堵缝材料；组分角色为材料学解释，不给出未公开配比',number=1)
for x,idx,title,body,c in [
    (.65,'01','苎麻','纤维填塞\n填充板缝、形成堵缝体',WOOD),
    (4.82,'02','桐油','油性组分\n参与复合堵缝材料',BLUE),
    (8.99,'03','石灰','无机组分\n与其他材料共同填缝',GREY),
]:
    panel(x,2.65,3.68,2.45)
    label(idx,x+.23,2.88,c,1,17)
    text(title,x+.23,3.33,3.15,.65,29,c,True)
    text(body,x+.23,4.08,3.15,.86,18)
line(1.0,5.67,12.25,5.67,GOLD,2)
for x,title in [(1.0,'选材'),(4.5,'填塞与压实'),(9.0,'服役与维护')]:
    shape(x,5.58,.16,.16,GOLD,MSO_SHAPE.OVAL)
    text(title,x+.28,5.39,2.7,.5,20)
text('守正创新：保存技艺，也用现代材料方法解释其有效条件。',.85,6.17,11.5,.4,17,BLUE)

# 4 — requirements
new('第二章  材料之困——现代巨轮如何抵御海洋侵蚀？','从“挡住水”，到同时回答三道材料题',
    '现代海洋防护的目标，是可靠性、可维护性与环境相容性的协同。',
    '来源 [5] IMO Anti-fouling systems；[7] Hao 等，Materials，2025；图为通用机理说明',dark=True,number=2)
cards=[(.65,'01','海水腐蚀','电解质 + 金属表面反应','需求：阻隔介质、维持附着'),
       (4.82,'02','生物污损','附着 → 生长 → 表面改变','需求：抑制附着或便于脱离'),
       (8.99,'03','绿色约束','施工、服役与维护的排放','需求：兼顾性能与环境影响')]
for x,num,title,mechanism,need in cards:
    panel(x,2.65,3.68,3.45,'11334B')
    label(num,x+.23,2.9,GOLD,1,17)
    text(title,x+.23,3.45,3.18,.6,28,WHITE,True)
    text(mechanism,x+.23,4.35,3.18,.95,19,MUTED)
    text(need,x+.23,5.34,3.18,.65,16,CYAN)
text('核心转变：从单一材料指标，走向服役条件下的系统性能。',.7,6.35,11.9,.4,18,WHITE)

# 5 — case evidence
new('第三章  科技破题——一层涂层背后的中国创新','核心案例：把防腐、防污与界面设计做成体系',
    '海洋防腐防污涂层体系关键技术及其应用 · 第一完成单位：集美大学',
    '来源 [2] 中国腐蚀与防护学会，2026-04-30，科技成果评价摘要 · /site/content/3294.html',number=3)
panel(.65,2.59,7.15,3.7)
label('评价摘要中的三组技术创新',.96,2.9,BLUE,6)
for yy,num,title in [(3.50,'01','高固低黏环氧 + 片状填料'),(4.32,'02','梯度化半互穿网络连接涂层'),(5.14,'03','有机硅结构调控 + 智能防污')]:
    label(num,1.0,yy,GOLD,.45,16)
    text(title,1.62,yy-.03,5.63,.52,23)
label('成果评价证书',8.25,2.89,BLUE,4.4)
text('中腐评价字〔2026〕\n第7号',8.25,3.50,4.35,1.1,24,INK,True)
text('摘要公开技术路线；\n完整配方与实验条件未披露。',8.25,4.72,4.35,.65,17,GREY)
shape(8.27,5.49,4.24,.65,'F5E8D3',MSO_SHAPE.ROUNDED_RECTANGLE)
text('以下自绘图均标明原理与事实边界',8.43,5.65,3.95,.32,14,'825A25',True)

# 6 — epoxy
new('第三章  科技破题  /  组成 → 结构 → 机理 → 性能','从树脂到填料：让腐蚀介质更难穿透',
    '项目路线：二聚酸改性高固低黏环氧 + 钛酸钾镁晶片 + 双电层片状填料',
    '来源 [2] 评价摘要第1项；[7] Hao 等，Materials，2025 · 图为一般屏障机理，非实测截面',number=3)
panel(.65,2.58,7.1,3.82)
label('水 / 氧 / 离子：可能的渗透路径',1.0,2.87,BLUE,6.2)
shape(1.0,3.55,6.4,1.65,'BCDCDD')
shape(1.0,5.22,6.4,.45,'8C9BA4')
for row in range(3):
    for col in range(8):
        xx,yy=1.3+col*.8,3.83+row*.48
        if col<7: line(xx,yy,xx+.8,yy+.12,BLUE,1.2)
        if row<2: line(xx,yy,xx+.13,yy+.48,BLUE,1.2)
        shape(xx-.045,yy-.045,.09,.09,BLUE,MSO_SHAPE.OVAL)
for xx,yy in [(1.55,4.1),(3.03,4.35),(4.37,3.89),(5.4,4.61)]:
    shape(xx,yy,.85,.08,GOLD)
for a,b in [((2.18,3.28),(2.18,3.88)),((2.18,3.88),(2.52,4.08)),((2.52,4.08),(2.34,4.54)),((2.34,4.54),(2.74,4.88))]:
    line(*a,*b,'CC8543',2.5)
label('固化网络 / 片状填料示意',1.0,5.86,BLUE,4.3,15)
label('钢基材',5.4,5.86,GREY,2,15)
for y,head,body in [(2.62,'组成设计','高固低黏树脂与片状填料协同'),(3.85,'结构与机理','固化网络与片状结构延缓介质迁移'),(5.08,'性能与边界','摘要报告高附着、长效防腐\n具体工况未披露')]:
    label(head,8.15,y,BLUE,4.4,16)
    text(body,8.15,y+.39,4.45,.78,20)

# 7 — silicone
new('第三章  科技破题  /  表面性质 → 附着 → 脱离','有机硅防污：同时面向动态与静态工况',
    '项目路线：功能树脂共聚；智能控释防污单元与 pH 响应双亲表面',
    '来源 [2] 评价摘要第3项；[8] Swain & Shinjo，IJMS，2014 · 图示为通用污损释放原理',number=3)
panel(.65,2.58,7.1,3.82)
label('污损释放：表面与水动力共同作用 · 原理示意',.98,2.88,BLUE,6.4)
shape(1.0,5.13,6.35,.48,'77C8CC')
shape(1.0,5.63,6.35,.25,'8C9BA4')
for x,y in [(2.05,4.67),(3.82,4.66),(5.67,3.97)]:
    shape(x,y,.48,.48,WOOD,MSO_SHAPE.OVAL)
    line(x+.24,y+.45,x+.24,5.1,WOOD,1)
for yy in [3.53,3.90]:
    line(1.14,yy,5.20,yy,BLUE,2)
    arrow(5.25,yy-.12,.42,BLUE)
line(5.95,4.35,6.35,3.73,GOLD,2)
text('附着较弱 ≠ 完全不附着',1.0,6.00,6.1,.34,17,BLUE,True)
for y,head,body in [(2.62,'动态场景','聚硼硅氧烷等功能树脂共聚\n提高高速运动时抗破损性'),(3.85,'静态场景','智能控释与响应表面\n改善静态防污期效'),(5.08,'解释边界','低表面能不等于不附着\n也不等于绝对无害')]:
    label(head,8.15,y,BLUE,4.4,16)
    text(body,8.15,y+.39,4.45,.78,20)

# 8 — stack
new('第三章  科技破题  /  体系配套与界面设计','多层涂层的短板，往往藏在界面',
    '从“每层性能好”，走向“整个体系在海洋环境中稳定工作”。',
    '来源 [2] 评价摘要第2项；层序为功能示意，不代表实际膜厚、配方或施工规范',dark=True,number=3)
layers=[(2.76,'有机硅防污层','表面性质 / 力学性能',CYAN),(3.48,'梯度半互穿网络连接层','固化动力学 / 层间结合',GOLD),(4.20,'环氧防腐层','成膜完整性 / 屏障',BLUE),(4.92,'钢基材','表面处理 / 清洁度','798D99')]
for y,name,role,c in layers:
    shape(.90,y,6.70,.52,c)
    text(name,1.08,y+.07,3.15,.38,17,NAVY if c in [CYAN,GOLD] else WHITE,True)
    line(7.61,y+.26,8.14,y+.26,c,1.2)
    text(role,8.32,y+.05,4.2,.44,20,WHITE)
text('常见失效风险',.93,5.83,3,.4,19,GOLD,True)
text('污染界面  /  重涂窗口不匹配  /  浸泡与温差引起的应力',3.95,5.84,8.35,.63,18,MUTED)

# 9 — translational
new('第三章  科技破题  /  工程应用与产学研协同','从实验室到船体，检验的是整个体系',
    '产学研协同的价值，在于把材料性能转化为可施工、可检验、可维护的方案。',
    '来源 [2] 成果评价摘要；[6] 集美大学官网，2026-10-23（图片：科研处；仅作课堂引用）',number=3)
photo('jmu-expo-2.jpg',.65,2.65,5.2,3.52)
text('集美大学海装展展台；非该涂层实船应用照片',.7,6.24,5.2,.42,12,GREY)
label('评价摘要确认：已实现产品应用转化',6.2,2.68,BLUE,6.3,16)
text('造船  /  核电  /  海工',6.2,3.23,6.2,.65,28,INK,True)
text('相关体系：CCS / DNV-GL 认证\n符合 NORSOK M-501 要求。\n未披露船名、应用面积和寿命数据。',6.2,4.04,6.25,1.33,19)
for x,title in [(6.2,'科研'),(8.43,'制造'),(10.66,'应用反馈')]:
    panel(x,5.77,1.82,.55,'E0EBEF')
    text(title,x+.08,5.89,1.65,.32,16,BLUE,True,PP_ALIGN.CENTER)
arrow(8.08,5.91,.23,BLUE);arrow(10.31,5.91,.23,BLUE)

# 10 — strategy chain
new('第四章  向海图强——材料创新背后的国家战略','一层涂层，如何成为海洋发展的支撑力？',
    '把科技自立自强与新质生产力，落到“技术—验证—产业能力”的具体转换中。',
    '来源 [2] 工程应用摘要；[3] 国家“十五五”纲要第三篇、第三十二章 · 图为本汇报的理论分析',dark=True,number=4)
items=[(.72,'国家需求','海洋装备\n可靠服役'),(3.88,'技术能力','材料 / 界面\n自主研究'),(7.04,'工程能力','生产 / 施工\n评价与维护'),(10.20,'发展质量','耐久 / 绿色\n产业升级')]
for x,head,body in items:
    panel(x,2.78,2.42,2.3,'11334B')
    label(head,x+.18,3.0,GOLD,2.05,16)
    text(body,x+.18,3.75,2.05,1.05,23,WHITE,True)
for x in [3.29,6.45,9.61]: arrow(x,3.72,.39,CYAN)
text('科技自立自强',.8,5.55,3.5,.48,23,CYAN,True)
text('掌握问题定义、技术路径与验证方法',4.46,5.57,8.05,.45,21,WHITE)
text('新质生产力',.8,6.12,3.5,.48,23,GOLD,True)
text('以创新推动生产方式、质量与效率改善',4.46,6.14,8.05,.45,21,WHITE)

# 11 — ecological
new('第四章  向海图强  /  习近平生态文明思想与绿色发展','绿色防护，要在全生命周期中算账',
    '以人与自然和谐共生为价值取向，把环境代价纳入材料研究的评价目标。',
    '来源 [4]《美丽上海建设“十五五”规划》；[5] IMO AFS · 不把材料名称当作环境安全证明',number=4)
rows=[(2.73,'施工阶段','高固体分路线','关注挥发性排放与施工暴露'),(3.84,'服役阶段','污损释放路线','关注释放物、磨损与清洗影响'),(4.95,'维护阶段','长效可靠防护','关注重涂、清洗频次与废物处理')]
for y,stage,tech,question in rows:
    panel(.65,y,12.02,.88)
    text(stage,.91,y+.22,2.1,.4,21,BLUE,True)
    text(tech,3.35,y+.22,3.15,.4,21)
    text(question,7.08,y+.24,5.1,.4,19,GREY)
text('研究结论应回答：相对什么方案？在什么工况？环境收益有何证据？',.76,6.24,11.9,.38,19,BLUE,True)

# 12 — plan/Shanghai
new('第四章  向海图强  /  “十五五”与上海','“十五五”与上海：把绿色要求变成材料任务',
    '从国家海洋科技部署，到上海造船、涂装行业的绿色转型需求。',
    '来源 [3] 国家纲要第三十二章；[4] 沪府发〔2026〕16号，第四部分（一）及第七部分（三）',dark=True,number=4)
for y,head,body in [(2.73,'国家部署','加强海洋科技创新，保护海洋生态'),(3.84,'上海要求','造船、涂装行业 VOCs 深度治理'),(4.95,'专业分析','高固体分 + 长效防护 + 环境相容')]:
    shape(.7,y,.08,.86,GOLD)
    text(head,.99,y+.16,2.25,.5,24,CYAN,True)
    text(body,3.51,y+.18,8.7,.5,24,WHITE)
text('上海材料研发的切入点：用工程验证连接绿色需求与成果转化。',.76,6.23,11.9,.4,20,GOLD)

# 13 — young researchers
new('第五章  青春担当——以材料之力回应时代之问','把“服务国家需求”，写成可验证的研究任务',
    '材料与化工研究生的担当，从一个明确问题、一组可信结果和一次有效协作开始。',
    '本页为专业行动建议，不构成已实施项目或已取得成果的描述',number=5)
for x,y,num,title,body in [
    (.65,2.57,'01','做深材料','连接组成与结构，解释性能来源'),
    (6.82,2.57,'02','做实界面','在浸泡、应力与施工条件下验证'),
    (.65,4.58,'03','做严评价','记录工况、对照、重复与失效'),
    (6.82,4.58,'04','做通转化','与制造和应用端共同定义问题'),
]:
    panel(x,y,5.85,1.76)
    text(num,x+.21,y+.25,.75,.55,25,GOLD,True)
    text(title,x+1.03,y+.22,4.5,.54,25,BLUE,True)
    text(body,x+1.03,y+.95,4.5,.55,18)

# 14 — conclusion
new('第五章  青春担当  /  回到起点','改变的是材料与工艺，不变的是面向真实问题',
    '从一条船缝，到一层涂层：以专业实践回应向海图强的时代要求。',
    '结语为本汇报综合表达；全部绘图为可编辑教学示意',dark=True,number=5)
ship(.95,2.92,4.25,WOOD,True)
ship(8.12,2.92,4.25,CYAN,False)
arrow(6.05,3.61,.7,GOLD)
text('传统：材料经验 + 工艺积累',.75,5.06,5.3,.5,22,WHITE,True)
text('现代：机理研究 + 系统验证',7.5,5.06,5.1,.5,22,WHITE,True)
text('求真  ·  协同  ·  创新  ·  责任',1.3,6.02,10.7,.5,27,GOLD,True,PP_ALIGN.CENTER)

# 15 — honest references
new('参考资料','核心来源与证据边界',
    '事实来自原文；机理图为自绘示意；理论联系与研究建议属于本汇报分析。',
    '完整题名、可点击链接、摘录与图片版权说明见 sources/参考资料与核验记录.md',dark=True)
entries=[
    ('1','非遗网：中国水密隔舱福船制造技艺','官方名录数据 · 项目 11847 · 传统材料与工艺','www.ihchina.cn/directory_list'),
    ('2','学会：海洋防腐防污涂层体系关键技术及其应用','2026-04-30 · 中腐评价字〔2026〕第7号','www.cscp.org.cn/site/content/3294.html'),
    ('3','国家“十五五”规划纲要','第三篇：科技自立自强；第三十二章：海洋','中国政府网 · 2026-03-13'),
    ('4','《美丽上海建设“十五五”规划》','沪府发〔2026〕16号 · 造船、涂装 VOCs 治理','上海市人民政府 · 2026-09-10'),
    ('5–8','IMO / 集美大学 / 两篇专业文献','法规与协同；环氧屏障、有机硅污损释放机理','Materials 2025；IJMS 2014 · 详见清单'),
]
for i,(code,title,status,url) in enumerate(entries):
    yy=2.52+i*.77
    label(code,.73,yy+.04,GOLD,.45,16)
    text(title,1.25,yy,5.48,.36,18,WHITE,True)
    text(status,1.25,yy+.38,5.48,.32,13,MUTED)
    box=text(url,7.0,yy+.10,5.45,.56,15,CYAN)
    index=i+1 if i<4 else 5
    box.text_frame.paragraphs[0].runs[0].hyperlink.address=SOURCES[index]
    box.text_frame.paragraphs[0].runs[0].font.color.rgb=rgb(CYAN)
    box.text_frame.paragraphs[0].runs[0].font.underline=False


def finalize_notes():
    """Speaker notes with source boundaries and a 900-second speaking plan."""
    final = [
        ('我们从一条不起眼的木船接缝出发，讨论一个贯穿古今的问题：材料怎样让船在海上可靠服役？今天沿着传统工艺、现代挑战、中国科研创新、国家战略与青年担当五章展开，从专业事实进入思政分析。', [1,2,3,4]),
        ('木船上有两种不同的问题：船板接缝可能渗水，而某处进水后又需要限制影响范围。艌缝针对前者，是接缝密封；水密隔舱针对后者，是结构分隔。非遗网的公开名录说明，一两个船舱偶然受损时，隔舱技术可以避免海水涌入其他舱室；同时明确把苎麻、桐油和石灰等列为板缝堵塞材料。我们不能把隔舱设计的功效全部归于一种填缝材料。图中只是这两种尺度的原理关系，不是福船历史实物的复原工程图。保护传统技艺，也要保护对技术细节的准确理解。', [1]),
        ('苎麻、桐油与石灰，在非遗资料中是共同用于木板缝堵塞的材料。用材料科学的语言看，纤维填塞、油性组分与无机组分，需要同施工压实相配合，才能形成接缝中的材料体系。这里的组分角色是教学解释，不是对历史配比和化学反应的复原，因此不编造传统配方或分子结构。非遗资料还描述了师傅头与工匠协作、师徒口传心授的传承方式。它提示我们，材料选择和工艺经验从来不是彼此分离的。守正创新，就是保存这些知识，再用现代方法研究其有效条件与改进空间。', [1]),
        ('木船的密封问题没有消失，但钢铁巨轮面对的是叠加的服役条件。首先，金属处于电解质环境，水、氧和离子可能通过涂层缺陷或渗透路径到达基材，腐蚀损伤要求材料维持屏障和附着。其次，生物附着与生长改变表面状态，涉及阻力、维护和服役表现。IMO介绍防污体系时也指出，污损会拖慢船舶、增加燃油消耗，但我们不引入未经统一工况核实的百分比。第三，防护方案本身也有环境影响。IMO的AFS公约针对有害防污体系，对作为杀生剂的有机锡化合物等进行管制，并纳入对cybutryne的控制。这里不展开全部法律适用条件，而强调研究方向的改变：不能只追求有效，还要评价环境相容性。由此形成三道题：怎样延缓腐蚀，怎样控制污损，怎样降低整体环境代价？这些要求可能相互制约，因此需要体系设计。', [5,7]),
        ('核心案例已在中国腐蚀与防护学会官网核实：项目题名为海洋防腐防污涂层体系关键技术及其应用，2026年4月30日公布评价摘要，第一完成单位是集美大学。摘要中的创新横跨树脂与填料、连接涂层和有机硅防污材料。我们依据摘要介绍技术路线，同时把通用机理和项目事实分开，不把未公开的实验条件、寿命或应用规模补出来。', [2]),
        ('先看高固体分环氧。评价摘要明确列出了二聚酸改性高固低黏环氧树脂、钛酸钾镁晶片以及双电层片状填料，并报告开发了五种超高固体分环氧防腐涂料。材料学上的关键，是把配方、施工和成膜连接起来。高固体分不能自动等同于零溶剂；低黏度也不是脱离施工条件的抽象优点，而关系到铺展、混合与成膜。一般环氧屏障原理可由专业综述支持：固化网络与分散良好的片状填料有助于阻碍腐蚀介质迁移。图中的网络和金色片层只是这个原理的示意，不是该项目真实微观结构。评价摘要报告了高附着与长效防腐等特性，却没有给出这里可以复述的完整测试条件。因此，我们强调设计逻辑，不把教学图当成显微证据，也不杜撰耐盐雾时长或电化学指标。', [2,7]),
        ('再看有机硅。左图解释的是污损释放型材料的一般机理：低表面能、低模量和平滑表面可降低部分海洋生物的附着强度，水动力或清洗再促成脱离。这个机制可由Swain和Shinjo的论文支持，但那篇论文的实验对象是养殖网，不能把其中效果直接移植成船体或本项目数据。项目评价摘要给出了更具体的两类路径：通过聚硼硅氧烷等功能树脂共聚，提高有机硅涂层的附着、强度和高速运动条件下的抗破损性能；通过海生物智能控释防污单元与pH响应双亲表面，提高静态防污期效。这意味着动态抗破损与静态防污需要分别考虑。我们不把所有路线都说成纯物理防污，也不因材料名称就宣称绝对无毒。', [2,8]),
        ('单层性能好，不意味着组合后就稳定。防腐层要维持牢固附着和完整屏障，防污层又有自己的表面与力学要求，因此连接界面是系统设计的重要位置。评价摘要明确提出，通过调控混合树脂涂料固化过程的动力学参数，制备梯度化半互穿网络结构高强连接涂层，提高与防腐层、防污层的结合力。我们可以把它理解为从固化过程组织界面结构的设计思路，但不能据此画出未公开的精确相组成或梯度曲线。左图仅表示各层功能，并不代表真实厚度。值得追问的是：层间污染、重涂窗口和长期服役应力，会不会成为短板？这也是从单项性能走向系统可靠性的研究训练。', [2]),
        ('评价摘要确认技术成果实现了产品应用转化，在造船、核电和海工等领域应用；同时称相关涂层体系通过CCS、DNV-GL等船级社认证，符合NORSOK M-501标准要求。这是成果走向工程的证据，但摘要没有公开具体船名、面积和服役寿命，我们不作补充推测。左图来自集美大学官网的2026海洋装备博览会报道，反映学校与行业对接的场景，不是本涂层项目的实船应用照。由此引出的产学研分析是：科研端解释机理，制造端控制工艺，用户端提供服役反馈，共同完成从材料到方案的转换。', [2,6]),
        ('思政分析不能止于给涂层贴上国家战略标签。中间的转换过程必须讲清：国家需要可靠海洋装备，装备需要基础材料，而材料还要经过制造、施工、认证和服役验证，才能形成稳定产业能力。国家十五五纲要第三篇强调高水平科技自立自强与新质生产力，也提出推动科技创新和产业创新深度融合；第三十二章部署海洋强国与海洋科技创新。放回案例，自立自强不是只得到一个配方，更是掌握问题定义、技术路径和验证方法；新质生产力也不等于实验室里多一个新名词，而要看技术是否进入生产过程、改善质量与效率。评价摘要确认了工程转化，使这一联系具有事实支点；但我们不能由摘要独立推出国产替代率或宏观经济贡献。这种既看到战略意义、又尊重证据尺度的分析，才体现研究生层次。', [2,3]),
        ('习近平生态文明思想在材料课题中不是最后补一句保护海洋，而应进入评价目标。上海美丽上海十五五规划明确以人与自然和谐共生为方向，并统筹降碳、减污、扩绿、增长。对应本案例，施工阶段应考察配方和工艺的挥发性排放，高固体分提供技术方向，但收益仍需实测；服役阶段，要区分污损释放、控释和响应表面，不把机制不同误写成没有任何环境风险；维护阶段，要把重涂、清洗、废物处理纳入分析。我们的价值判断因此转化为三个可检验的问题：相对什么基准，在什么工况，系统边界到哪里？只有避免环境负担转移，才可能说明技术如何同时支持开发利用和生态保护。', [4,5]),
        ('十五五的联系可以分为两层原文和一层专业分析。国家纲要第三十二章明确提出加强海洋科技创新、巩固提升海洋装备制造业优势，并在下一节部署海洋生态环境保护。上海2026年9月公布的美丽上海建设十五五规划，明确要求造船、涂装等行业开展VOCs深度治理，并支持科技成果孵化转化。这使我们能够提出一个有依据的专业切入点：上海海洋相关装备场景需要更可靠、更绿色、能够稳定施工的材料，高固体分防腐体系、耐久界面和环境相容性研究因此具有需求牵引。但要注意，最后一层是我们从政策和案例推导出的材料研究建议，不是说这份生态环境规划设立了一个同名海洋新材料专项，也不声称该集美大学成果已经在上海落地。国家方向、地方需求与具体课题通过机制相连，而不是靠地名或口号相连。', [3,4]),
        ('青年担当需要落实到日常研究。第一，做深材料，不只报告某个性能提高，还要说明组成和结构怎样影响机制，证据是否支持解释。第二，做实界面，把浸泡、应力和施工条件纳入验证，不局限在理想小试片。第三，做严评价，保留工况、对照、重复与异常，既报告成功，也分析失效。第四，做通转化，与制造和应用端共同定义问题，理解工艺放大与实际维护的限制。传统工匠的经验值得尊重，现代科研则要求把认识转化为可复现、可交流的证据。服务国家需求，不是给每项工作加上宏大标题，而是把真实问题研究透，把每个结论做可信，让解决方案更接近海上服役。材料与化工的交叉优势，正可以体现在这一连续过程中。', [1,2,3]),
        ('最后回到一条船缝。古代工匠用材料经验和工艺积累解决渗漏，现代研究者用机理研究、界面设计和工程验证回应更复杂的海洋环境。改变的是材料与工艺，不变的是面向真实问题的智慧与担当。从填补木船的一条缝，到守护现代巨轮的一层涂层，我们愿以可信的材料研究，为海洋强国建设提供扎实的专业支撑。', [1,2,3]),
        ('参考页不计入讲述。完整来源、访问方式、核验摘录和图片来源见随附清单。图中箭头、网格和层序为自绘原理示意；科研项目的公开评价摘要不等于完整实验报告；政策对应和青年行动属于本汇报分析。', [1,2,3,4,5,6,7,8]),
    ]
    for i,(narrative,ids) in enumerate(final):
        n=dict(page=i+1,seconds=TIMES[i])
        NOTES.append(n)
        n['narrative']=narrative
        n['sources']=[SOURCES[j] for j in ids]
        n['status']='来源已读取；项目事实、通用机理与本汇报分析分别表述。原文未公开的信息不作推断。'
        prs.slides[i].notes_slide.notes_text_frame.text=(
            f"建议时长：{n['seconds']} 秒\n\n{narrative}\n\n"
            f"证据边界：{n['status']}\n来源：\n"+'\n'.join(n['sources']))


def verify():
    issues=[]
    for item in LAYOUT:
        if item['x'] < -.001 or item['y'] < -.001 or item['x']+item['w'] > W+.001 or item['y']+item['h'] > H+.001:
            issues.append({'kind':'bounds',**item})
    # Native backgrounds and diagram containers intentionally surround content.
    # Check text-to-text overlaps; small numerical tolerances avoid touching edges.
    for i,a in enumerate(LAYOUT):
        for b in LAYOUT[i+1:]:
            if a['slide']!=b['slide'] or not a['check'] or not b['check']: continue
            dx=min(a['x']+a['w'],b['x']+b['w'])-max(a['x'],b['x'])
            dy=min(a['y']+a['h'],b['y']+b['h'])-max(a['y'],b['y'])
            if dx>.02 and dy>.02:
                issues.append(dict(kind='text_overlap',slide=a['slide'],a=a['text'],b=b['text'],dx=dx,dy=dy))
    assert len(prs.slides)==15
    assert sum(n['seconds'] for n in NOTES)==900
    for part in prs.part.package.iter_parts():
        if part.content_type=='application/vnd.openxmlformats-officedocument.theme+xml':
            theme=etree.fromstring(part.blob)
            ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
            for node in theme.xpath('//a:clrScheme/a:hlink/a:srgbClr | //a:clrScheme/a:folHlink/a:srgbClr',namespaces=ns):
                node.set('val',CYAN)
            part._blob=etree.tostring(theme,xml_declaration=True,encoding='UTF-8',standalone=True)
    prs.save(OUT)
    with ZipFile(OUT) as z:
        assert z.testzip() is None
        for name in z.namelist():
            if name.endswith(('.xml','.rels')): etree.fromstring(z.read(name))
    loaded=Presentation(OUT)
    assert len(loaded.slides)==15
    textboxes=sum(s.has_text_frame and bool(s.text.strip()) for sl in loaded.slides for s in sl.shapes)
    pictures=sum(s.shape_type==13 for sl in loaded.slides for s in sl.shapes)
    for page,sl in enumerate(loaded.slides,1):
        for s in sl.shapes:
            assert s.left >= 0 and s.top >= 0 and s.left+s.width <= loaded.slide_width+2 and s.top+s.height <= loaded.slide_height+2, (page,s.name,'shape outside slide')
    report=dict(slides=15,seconds=900,editable_text_shapes=textboxes,pictures=pictures,
                zip_crc='passed',xml_parse='passed',pptx_reopen='passed',all_shape_bounds='passed',
                font=FONT,minimum_text_size=min(x['size'] for x in LAYOUT),
                text_layout_issues=issues,content_status='课堂汇报稿：核心来源已读取，证据边界见逐页备注与来源清单')
    (BASE/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
    (BASE/'layout.json').write_text(json.dumps(LAYOUT,ensure_ascii=False,indent=2))
    md=['# 15 分钟讲稿与证据边界\n','总时长：900 秒；参考资料页不计时。\n']
    for n in NOTES:
        md += [f"## 第 {n['page']} 页 · {n['seconds']} 秒\n",n['narrative']+'\n',f"核验状态：{n['status']}\n"]
        md += [f'- {url}' for url in n['sources']]
        md += ['\n']
    (BASE/'讲稿与备注.md').write_text('\n'.join(md))
    print(json.dumps(report,ensure_ascii=False,indent=2))
    assert not issues, 'Layout issues; inspect validation.json'


if __name__=='__main__':
    finalize_notes()
    verify()

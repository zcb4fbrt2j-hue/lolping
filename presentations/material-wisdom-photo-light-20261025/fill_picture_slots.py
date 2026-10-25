#!/usr/bin/env python3
"""Insert user photos into native picture slots, saving a NEW PPTX only."""
import argparse,hashlib,json
from pathlib import Path
from PIL import Image
from pptx import Presentation

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--input',type=Path,required=True)
    p.add_argument('--manifest',type=Path,required=True)
    p.add_argument('--output',type=Path,required=True)
    a=p.parse_args()
    if a.input.resolve()==a.output.resolve() or a.output.exists():raise ValueError('Output must be a NEW file; input is never overwritten')
    prs=Presentation(a.input);slots={s.name:(sl,s) for sl in prs.slides for s in sl.placeholders}
    rows=json.loads(a.manifest.read_text())['photos'];assert rows,'photos must not be empty'
    hashes={hashlib.sha256(s.image.blob).hexdigest() for sl in prs.slides for s in sl.shapes if hasattr(s,'image')}
    used=set();record=[]
    for row in rows:
        for key in ['slot','file','content','location','date','photographer','usage_permission']:
            if not row.get(key) or '请填写' in row[key]:raise ValueError('Complete metadata: '+key)
        name=row['slot']
        if name in used or name not in slots:raise ValueError('Unknown or repeated slot: '+name)
        used.add(name);sl,ph=slots[name]
        if int(ph.placeholder_format.type)!=18 or not hasattr(ph,'insert_picture'):raise ValueError('Slot is already filled: '+name)
        f=Path(row['file'])
        if not f.is_absolute():f=a.manifest.resolve().parent/f
        blob=f.read_bytes();digest=hashlib.sha256(blob).hexdigest()
        if digest in hashes:raise ValueError('Repeated image is forbidden, including cover/body: '+name)
        hashes.add(digest)
        with Image.open(f) as im:im.verify()
        with Image.open(f) as im:
            if im.width<1920:raise ValueError('Use original width >=1920px; no upscaling: '+str(f))
        box=(ph.left,ph.top,ph.width,ph.height);pic=ph.insert_picture(str(f))
        assert box==(pic.left,pic.top,pic.width,pic.height)
        for s in list(sl.shapes):
            if s.name==name+'_LABEL':s._element.getparent().remove(s._element)
        sl.notes_slide.notes_text_frame.text+='\n用户补充照片：'+json.dumps({**row,'sha256':digest},ensure_ascii=False)
        record.append({**row,'sha256':digest})
    a.output.parent.mkdir(parents=True,exist_ok=True);prs.save(a.output)
    Presentation(a.output)
    print('Saved NEW presentation:',a.output)
    print(json.dumps(record,ensure_ascii=False,indent=2))
    print('Next: check identity, crop and captions manually; update pending observation text, then render and review.')

if __name__=='__main__':main()

"""Validate the editable sample delivery, including native photo crop ratios."""
from pathlib import Path
from zipfile import ZipFile
import json, hashlib, io
from lxml import etree
from pptx import Presentation
from PIL import Image
import fitz
ROOT=Path(__file__).resolve().parent
PPTX=ROOT/'build/material-wisdom-v2-samples.pptx'
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
rows=[]
with ZipFile(PPTX) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith('.xml'): etree.fromstring(z.read(name))
pres=Presentation(PPTX)
assert len(pres.slides)==3
assert abs(pres.slide_width/pres.slide_height-16/9)<.001
for i,slide in enumerate(pres.slides,1):
    texts=[s for s in slide.shapes if s.has_text_frame and s.text.strip()]
    pics=[s for s in slide.shapes if s.shape_type==13]
    assert len(texts)>=7 and pics
    for shape in pics:
        with Image.open(io.BytesIO(shape.image.blob)) as im:
            ratio=im.width*(1-shape.crop_left-shape.crop_right)/(im.height*(1-shape.crop_top-shape.crop_bottom))
        assert abs(ratio/(shape.width/shape.height)-1)<.001, (i,'stretched photo')
    rows.append({'slide':i,'editable_text_objects':len(texts),'picture_objects':len(pics),'has_notes':slide.has_notes_slide})
previews=sorted((ROOT/'build/previews').glob('slide-*.png'))
assert len(previews)==3
for file in previews:
    with Image.open(file) as im:
        assert im.width >= 1500 and abs(im.width/im.height-16/9)<.003
        im.verify()
with fitz.open(ROOT/'build/material-wisdom-v2-samples.pdf') as pdf:
    assert len(pdf)==3
    assert all(len(page.get_text().strip())>40 for page in pdf)
report={'pptx_zip_xml_valid':True,'slide_count':3,'aspect':'16:9','photos_preserve_aspect':True,'pdf_pages':3,'preview_count':3,'slides':rows,'files':[]}
for file in [PPTX,ROOT/'build/material-wisdom-v2-samples.pdf',*previews]:
    report['files'].append({'file':str(file.relative_to(ROOT)),'bytes':file.stat().st_size,'sha256':hashlib.sha256(file.read_bytes()).hexdigest()})
(ROOT/'build/integrity.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(report,ensure_ascii=False,indent=2))
